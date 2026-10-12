"""Tool đọc/ghi topcv_jobs.xlsx cho agent LangChain / LangGraph.

Tools: write_excel, find_job, read_pending_apply, mark_ai_applied
Phụ thuộc: pip install openpyxl pydantic langchain-core
"""
import json
import os
import unicodedata
from pathlib import Path
from typing import Literal

from langchain_core.tools import tool
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from pydantic import BaseModel, Field
from dotenv import load_dotenv
ROOT = Path(__file__).resolve().parents[1]          # D:\agent_auto_job
load_dotenv(ROOT / ".env")
EXCEL_PATH = Path(os.getenv("TOPCV_EXCEL_PATH", ROOT / "topcv_jobs.xlsx"))
SHEET = "Jobs"

# (tên field, tiêu đề cột Excel) — STT là cột đầu, tự đánh số
COLUMNS = [
    ("ten_viec_lam", "Tên việc làm"),
    ("cong_ty", "Công ty"),
    ("link", "Link"),
    ("muc_luong", "Mức lương"),
    ("dia_diem", "Địa điểm"),
    ("kinh_nghiem", "Kinh nghiệm"),
    ("cap_bac_hinh_thuc", "Cấp bậc / Hình thức"),
    ("tom_tat_mo_ta", "Tóm tắt mô tả công việc"),
    ("tom_tat_yeu_cau", "Tóm tắt yêu cầu ứng viên"),
    ("tom_tat_quyen_loi", "Tóm tắt quyền lợi"),
    ("dia_diem_thoi_gian", "Địa điểm & thời gian làm việc"),
    ("han_nop_thong_tin_khac", "Hạn nộp / Thông tin khác"),
    ("trang", "Trang"),
    ("human_apply", "human_apply"),
    ("ai_apply", "ai_apply"),
]
HEADERS = ["STT"] + [h for _, h in COLUMNS]
FIELDS = [f for f, _ in COLUMNS]
BRIEF_FIELDS = ["stt", "ten_viec_lam", "cong_ty", "link", "trang", "human_apply", "ai_apply"]
WRAP_FIELDS = {"tom_tat_mo_ta", "tom_tat_yeu_cau", "tom_tat_quyen_loi", "dia_diem_thoi_gian"}


class JobRow(BaseModel):
    """Một dòng job trong topcv_jobs.xlsx (STT tự đánh số)."""

    ten_viec_lam: str = Field(description="Tên việc làm / tiêu đề job")
    cong_ty: str = Field(description="Tên công ty")
    link: str = Field(description="URL trang chi tiết job (khóa chống trùng)")
    muc_luong: str = Field(description='Mức lương, vd "2 - 3 triệu", "Thoả thuận"')
    dia_diem: str = Field(description="Địa chỉ làm việc chi tiết")
    kinh_nghiem: str = Field(description="Yêu cầu kinh nghiệm ghi trên tin")
    cap_bac_hinh_thuc: str = Field(description="Cấp bậc / hình thức: Thực tập, Nhân viên, Full-time...")
    tom_tat_mo_ta: str = Field(description="Tóm tắt mô tả công việc")
    tom_tat_yeu_cau: str = Field(description="Tóm tắt yêu cầu ứng viên: kỹ năng, học vấn, công nghệ")
    tom_tat_quyen_loi: str = Field(description="Tóm tắt quyền lợi: lương thưởng, đào tạo, phúc lợi")
    dia_diem_thoi_gian: str = Field(description="Địa điểm & thời gian làm việc")
    han_nop_thong_tin_khac: str = Field(description="Hạn nộp hồ sơ, số lượng tuyển, thông tin khác")
    trang: int = Field(ge=1, description="Số trang danh sách nơi tìm thấy job")
    human_apply: Literal[0, 1] = Field(default=0, description="Người dùng tự điền; khi ghi dòng mới để 0")
    ai_apply: Literal[0, 1] = Field(default=0, description="Agent đã apply chưa; khi ghi dòng mới để 0")


def _norm_link(link) -> str:
    return str(link or "").strip().split("?")[0].rstrip("/").lower()


def _flag(v) -> int:
    return 1 if str(v).strip() in {"1", "1.0", "true", "True"} else 0


def _init_sheet(ws) -> None:
    """Ghi header vào sheet trống. Phải gọi trước khi chạm vào bất kỳ ô nào
    (truy cập ô làm openpyxl coi dòng 1 đã dùng và append xuống dòng 2)."""
    ws.append(HEADERS)
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.freeze_panes = "A2"
    for i, h in enumerate(HEADERS, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = 45 if h.startswith("Tóm tắt") else 22


def _open():
    if EXCEL_PATH.exists():
        wb = load_workbook(EXCEL_PATH)
        if SHEET in wb.sheetnames:
            return wb, wb[SHEET]
        ws = wb.create_sheet(SHEET)
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = SHEET
    _init_sheet(ws)
    return wb, ws


def _rows(ws) -> list[dict]:
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not any(r):
            continue
        d = dict(zip(["stt"] + FIELDS, r))
        d["human_apply"] = _flag(d.get("human_apply"))
        d["ai_apply"] = _flag(d.get("ai_apply"))
        rows.append(d)
    return rows


def _save(wb) -> str | None:
    try:
        wb.save(EXCEL_PATH)
    except PermissionError:
        return "LỖI: không ghi được file, có thể đang mở trong Excel. Đóng file rồi gọi lại."
    return None


@tool("write_excel", args_schema=JobRow)
def write_excel(**fields) -> str:
    """Append 1 job vào topcv_jobs.xlsx. Gọi NGAY sau khi tóm tắt xong mỗi job hợp lệ.
    Bỏ qua nếu link đã tồn tại (không ghi đè, giữ nguyên human_apply)."""
    job = JobRow(**fields)
    wb, ws = _open()
    key = _norm_link(job.link)
    existing = _rows(ws)
    for r in existing:
        if _norm_link(r["link"]) == key:
            return f"BỎ QUA: job đã có trong Excel (STT {r['stt']}): {r['ten_viec_lam']} - {r['cong_ty']}"
    stt = len(existing) + 1
    ws.append([stt] + [getattr(job, f) for f in FIELDS])
    for col, (field, _) in enumerate(COLUMNS, start=2):
        if field in WRAP_FIELDS:
            ws.cell(ws.max_row, col).alignment = Alignment(wrap_text=True, vertical="top")
    err = _save(wb)
    return err or f"ĐÃ GHI: STT {stt} - {job.ten_viec_lam} - {job.cong_ty}"


def _fold(s) -> str:
    """Hạ chữ thường + bỏ dấu tiếng Việt để search không phân biệt dấu."""
    s = str(s or "").lower().replace("đ", "d")
    return "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")


def _brief(r: dict) -> dict:
    return {k: r[k] for k in BRIEF_FIELDS}


@tool("find_job")
def find_job(query: str) -> str:
    """Kiểm tra job đã có trong topcv_jobs.xlsx chưa (tránh làm trùng).
    query: tên việc làm + tên công ty, hoặc link (không phân biệt hoa/thường và dấu).
    Nhiều từ → dòng phải chứa TẤT CẢ các từ (tìm trong tên + công ty + link).
    Trả "ĐÃ CÓ: ..." (bỏ qua job này) hoặc "CHƯA CÓ" (làm tiếp)."""
    if not EXCEL_PATH.exists():
        return "CHƯA CÓ"
    tokens = [_norm_link(t) if t.lower().startswith("http") else _fold(t) for t in query.split()]
    tokens = [t for t in tokens if t]
    if not tokens:
        return "LỖI: query rỗng, nhập tên việc làm + công ty hoặc link"
    _, ws = _open()
    hits = [
        r for r in _rows(ws)
        if all(t in _fold(f"{r['ten_viec_lam']} {r['cong_ty']} {r['link']}") for t in tokens)
    ]
    if not hits:
        return "CHƯA CÓ"
    return "ĐÃ CÓ: " + "; ".join(f"STT {r['stt']} - {r['ten_viec_lam']} - {r['cong_ty']}" for r in hits[:3])


@tool("read_pending_apply")
def read_pending_apply() -> str:
    """Danh sách job cần apply: human_apply=1 và ai_apply=0 (dùng cho skill auto-apply).
    Trả JSON gồm stt, tên, công ty, link, trang, human_apply, ai_apply."""
    if not EXCEL_PATH.exists():
        return json.dumps({"total": 0, "rows": []})
    _, ws = _open()
    rows = [r for r in _rows(ws) if r["human_apply"] == 1 and r["ai_apply"] == 0]
    return json.dumps({"total": len(rows), "rows": [_brief(r) for r in rows]}, ensure_ascii=False, default=str)


@tool("mark_ai_applied")
def mark_ai_applied(link: str) -> str:
    """Đặt ai_apply=1 cho dòng có link này. CHỈ gọi sau khi đã thấy xác nhận nộp đơn thành công."""
    wb, ws = _open()
    link_idx = HEADERS.index("Link")
    ai_idx = HEADERS.index("ai_apply")
    key = _norm_link(link)
    for row in ws.iter_rows(min_row=2):
        if _norm_link(row[link_idx].value) == key:
            row[ai_idx].value = 1
            err = _save(wb)
            return err or f"ĐÃ CẬP NHẬT: ai_apply=1 cho STT {row[0].value}"
    return "KHÔNG TÌM THẤY: link không có trong Excel"


TOOLS = [write_excel, find_job, read_pending_apply, mark_ai_applied]
